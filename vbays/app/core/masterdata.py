"""Import master data from Excel (.xlsx) or CSV files.

Each dataset has a fixed list of columns (see DATASETS). The import is done in
two steps: preview (find every mistake, change nothing) then confirm (save).
Existing rows with the same key are updated; new rows are added.
Money columns are in RUPEES in the file and stored as paise.
"""
import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core import models as m
from app.core.permissions import ROLES


@dataclass
class Col:
    name: str
    kind: str = "str"  # str | int | money | decimal | bool | role
    required: bool = True
    attr: str | None = None  # model attribute (defaults to name)
    help: str = ""


@dataclass
class Dataset:
    key: str
    title: str
    model: type
    key_cols: list[str]
    cols: list[Col]
    description: str = ""
    sample_rows: list[dict] = field(default_factory=list)

    @property
    def filename(self) -> str:
        return f"{self.key}.csv"


DATASETS: dict[str, Dataset] = {
    d.key: d
    for d in [
        Dataset(
            "rate_card", "Rate card", m.RateCardItem, ["item_code"],
            [
                Col("item_code"), Col("item_name"), Col("category"), Col("unit", help="sqft / rft / nos / set"),
                Col("basic_rate", "money", attr="basic_rate_paise"),
                Col("premium_rate", "money", attr="premium_rate_paise"),
                Col("luxury_rate", "money", attr="luxury_rate_paise"),
                Col("gst_percent", "decimal"), Col("notes", required=False),
            ],
            "Selling rates used for estimates and quotations (₹, before GST).",
        ),
        Dataset(
            "materials_finishes", "Materials & finishes", m.Material, ["code"],
            [
                Col("code"), Col("name"), Col("type", help="board / laminate / acrylic / pu / edge_band / adhesive / other"),
                Col("brand", required=False), Col("thickness_mm", "decimal", required=False),
                Col("sheet_length_mm", "int", required=False), Col("sheet_width_mm", "int", required=False),
                Col("unit"), Col("cost_rate", "money", attr="cost_rate_paise"),
                Col("sell_rate", "money", required=False, attr="sell_rate_paise"),
            ],
        ),
        Dataset(
            "hardware", "Hardware", m.HardwareItem, ["code"],
            [
                Col("code"), Col("name"), Col("brand", required=False), Col("type"), Col("unit"),
                Col("cost_rate", "money", attr="cost_rate_paise"),
                Col("sell_rate", "money", required=False, attr="sell_rate_paise"),
            ],
        ),
        Dataset(
            "product_catalog", "Product catalog (standard units)", m.ProductCatalogItem, ["code"],
            [
                Col("code"), Col("name"), Col("category"),
                Col("std_width_mm", "int", required=False), Col("std_height_mm", "int", required=False),
                Col("std_depth_mm", "int", required=False),
                Col("default_board_code", required=False), Col("default_finish_code", required=False),
                Col("default_hardware", required=False, help="CODE:QTY;CODE:QTY"),
            ],
        ),
        Dataset(
            "vendors", "Vendors", m.Vendor, ["name"],
            [
                Col("name"), Col("contact_person", required=False), Col("phone", required=False),
                Col("email", required=False), Col("gstin", required=False),
                Col("items_supplied", required=False), Col("payment_terms", required=False),
            ],
        ),
        Dataset(
            "staff_roles", "Staff & roles", m.User, ["phone"],
            [
                Col("name"), Col("role", "role"), Col("phone"),
                Col("email", required=False), Col("telegram_username", required=False),
            ],
            "Adds staff logins. New staff get no password; the owner sets one in Users, "
            "or they use Telegram only.",
        ),
        Dataset(
            "production_stages", "Factory production stages", m.ProductionStage, ["sequence"],
            [
                Col("sequence", "int"), Col("name"), Col("planned_hours", "decimal"),
                Col("needs_photo", "bool"), Col("needs_qc", "bool"),
            ],
        ),
        Dataset(
            "payment_milestones", "Payment milestone templates", m.PaymentMilestoneTemplate,
            ["template_name", "sequence"],
            [
                Col("template_name"), Col("sequence", "int"), Col("milestone_name"),
                Col("percent", "decimal"), Col("trigger", help="booking / before_production / before_installation / handover"),
            ],
            "Percentages in each template must add up to 100.",
        ),
        Dataset(
            "whatsapp_templates", "WhatsApp message templates", m.MessageTemplate, ["name", "language"],
            [
                Col("name"), Col("language", help="te / en / te-en"),
                Col("category", help="utility / marketing / authentication"),
                Col("body", help="Use {{1}}, {{2}} … for customer name, date etc."),
                Col("meta_approved", "bool", required=False),
            ],
            "Templates must also be approved by Meta before WhatsApp will deliver them.",
        ),
    ]
}


@dataclass
class ImportResult:
    rows: list[dict]
    errors: list[str]
    to_add: int = 0
    to_update: int = 0

    @property
    def ok(self) -> bool:
        return not self.errors and bool(self.rows)


def read_table(filename: str, content: bytes) -> list[dict]:
    """Read .csv or .xlsx into a list of {column: text}."""
    if filename.lower().endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip().lower() if h is not None else "" for h in next(it, [])]
        rows = []
        for values in it:
            if values is None or all(v is None or str(v).strip() == "" for v in values):
                continue
            rows.append({header[i]: ("" if v is None else str(v)).strip() for i, v in enumerate(values) if i < len(header)})
        return rows
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [
        {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        for row in reader
        if any((v or "").strip() for v in row.values())
    ]


_TRUE = {"y", "yes", "true", "1"}
_FALSE = {"n", "no", "false", "0", ""}


def _convert(col: Col, raw: str):
    raw = raw.replace("₹", "").strip() if col.kind in ("money", "decimal", "int") else raw.strip()
    if raw == "":
        return None
    if col.kind == "str":
        return raw
    if col.kind == "role":
        key = raw.lower().replace(" ", "_")
        if key not in ROLES:
            raise ValueError(f"unknown role '{raw}' (use one of: {', '.join(ROLES)})")
        return key
    if col.kind == "bool":
        v = raw.lower()
        if v in _TRUE:
            return True
        if v in _FALSE:
            return False
        raise ValueError(f"'{raw}' is not Y or N")
    try:
        num = Decimal(raw.replace(",", ""))
    except InvalidOperation:
        raise ValueError(f"'{raw}' is not a number") from None
    if num < 0:
        raise ValueError("cannot be negative")
    if col.kind == "int":
        if num != num.to_integral_value():
            raise ValueError(f"'{raw}' must be a whole number")
        return int(num)
    if col.kind == "money":
        return int((num * 100).quantize(Decimal("1")))
    return num  # decimal


def validate(ds: Dataset, raw_rows: list[dict]) -> ImportResult:
    errors: list[str] = []
    rows: list[dict] = []
    if not raw_rows:
        return ImportResult([], ["The file has no data rows."])
    missing = [c.name for c in ds.cols if c.required and c.name not in raw_rows[0]]
    if missing:
        return ImportResult([], [f"Missing column(s): {', '.join(missing)}"])

    seen = set()
    for i, raw in enumerate(raw_rows, start=2):  # row 1 is the header
        row, row_ok = {}, True
        for c in ds.cols:
            try:
                val = _convert(c, raw.get(c.name, ""))
            except ValueError as exc:
                errors.append(f"Row {i}, {c.name}: {exc}")
                row_ok = False
                continue
            if val is None and c.required:
                errors.append(f"Row {i}, {c.name}: is empty")
                row_ok = False
            row[c.attr or c.name] = val
        if not row_ok:
            continue
        key = tuple(row[k] for k in ds.key_cols)
        if key in seen:
            errors.append(f"Row {i}: duplicate {', '.join(ds.key_cols)} {key}")
            continue
        seen.add(key)
        rows.append(row)

    if ds.key == "payment_milestones" and not errors:
        totals: dict[str, Decimal] = {}
        for r in rows:
            totals[r["template_name"]] = totals.get(r["template_name"], Decimal(0)) + r["percent"]
        for name, total in totals.items():
            if total != 100:
                errors.append(f"Template '{name}': percentages add up to {total}, not 100")
    if ds.key == "staff_roles":
        for r in rows:
            if r.get("email"):
                r["email"] = r["email"].lower()
    return ImportResult(rows, errors)


def _find(db: Session, ds: Dataset, row: dict):
    q = select(ds.model)
    for k in ds.key_cols:
        q = q.where(getattr(ds.model, k) == row[k])
    return db.scalar(q)


def preview(db: Session, ds: Dataset, raw_rows: list[dict]) -> ImportResult:
    res = validate(ds, raw_rows)
    if res.errors:
        return res
    for r in res.rows:
        if _find(db, ds, r):
            res.to_update += 1
        else:
            res.to_add += 1
    return res


def apply(db: Session, ds: Dataset, res: ImportResult) -> tuple[int, int]:
    if not res.ok:
        raise ValueError("Cannot import a file with errors.")
    added = updated = 0
    for r in res.rows:
        obj = _find(db, ds, r)
        if obj is None:
            obj = ds.model(**r)
            db.add(obj)
            added += 1
        else:
            for k, v in r.items():
                setattr(obj, k, v)
            updated += 1
    db.flush()
    return added, updated


def count(db: Session, ds: Dataset) -> int:
    return db.scalar(select(func.count()).select_from(ds.model)) or 0


def template_csv(ds: Dataset) -> str:
    return ",".join(c.name for c in ds.cols) + "\n"
